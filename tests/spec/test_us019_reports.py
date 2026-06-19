"""US-019 — specs/features/epic-07-dashboard.md"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import load_workbook

from ppe_detection.dashboard.reports import export_report, generate_summary
from ppe_detection.storage.occurrence_store import OccurrenceRecord, OccurrenceStore


@pytest.mark.spec("US-019")
class TestUS019Reports:
    @pytest.fixture
    def store(self, tmp_path):
        s = OccurrenceStore(tmp_path / "test.db")
        s.record(OccurrenceRecord(
            id=None, camera_id="cam1", track_id=1,
            reason="Sem EPI: Capacete", missing_ppe=["helmet"],
            created_at=datetime.now(),
        ))
        return s

    def test_ac_019_1_summary_indicators(self, store):
        summary = generate_summary(store)
        assert summary["total_infractions"] == 1
        assert "compliance_rate" in summary

    def test_ac_019_2_export_xlsx(self, store, tmp_path):
        path = tmp_path / "report.xlsx"
        export_report(store, path, "xlsx")
        wb = load_workbook(path)
        assert "Resumo" in wb.sheetnames
        assert "Ocorrências" in wb.sheetnames

    def test_ac_019_3_export_pdf(self, store, tmp_path):
        path = tmp_path / "report.pdf"
        export_report(store, path, "pdf")
        assert path.exists()
        assert path.read_bytes()[:4] == b"%PDF"
