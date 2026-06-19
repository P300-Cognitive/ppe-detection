"""Geração de relatórios — spec: specs/features/epic-07-dashboard.md (US-019)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from fpdf import FPDF
from openpyxl import Workbook

from ppe_detection.storage.occurrence_store import OccurrenceStore


def generate_summary(
    store: OccurrenceStore,
    from_date: datetime | None = None,
    to_date: datetime | None = None,
    camera_id: str | None = None,
) -> dict:
    return store.count_by_period(from_date, to_date, camera_id)


def export_report(
    store: OccurrenceStore,
    output_path: Path,
    fmt: str,
    from_date: datetime | None = None,
    to_date: datetime | None = None,
    camera_id: str | None = None,
) -> Path:
    summary = generate_summary(store, from_date, to_date, camera_id)
    occurrences = store.search(from_date, to_date, camera_id, limit=10_000)

    if fmt == "xlsx":
        return _export_xlsx(output_path, summary, occurrences)
    if fmt == "pdf":
        return _export_pdf(output_path, summary, occurrences)
    raise ValueError(f"Formato não suportado: {fmt}")


def _export_xlsx(path: Path, summary: dict, occurrences) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Indicador", "Valor"])
    ws.append(["Total de infrações", summary["total_infractions"]])
    ws.append(["Taxa de conformidade", summary["compliance_rate"]])

    ws2 = wb.create_sheet("Ocorrências")
    ws2.append(["ID", "Câmera", "Track", "Motivo", "EPIs ausentes", "Data"])
    for occ in occurrences:
        ws2.append([
            occ.id, occ.camera_id, occ.track_id, occ.reason,
            ", ".join(occ.missing_ppe), occ.created_at.isoformat(),
        ])

    wb.save(path)
    return path


def _export_pdf(path: Path, summary: dict, occurrences) -> Path:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text="Relatório de Conformidade EPI")
    pdf.ln(10)
    pdf.set_font("Helvetica", size=11)
    pdf.cell(text=f"Total de infrações: {summary['total_infractions']}")
    pdf.ln(6)
    pdf.cell(text=f"Taxa de conformidade: {summary['compliance_rate']:.1%}")
    pdf.ln(10)

    pdf.set_font("Helvetica", size=10)
    for occ in occurrences[:50]:
        line = (
            f"#{occ.id} [{occ.camera_id}] {occ.created_at.strftime('%d/%m/%Y %H:%M')} "
            f"- {occ.reason}"
        )
        pdf.cell(text=line)
        pdf.ln(5)

    pdf.output(str(path))
    return path
